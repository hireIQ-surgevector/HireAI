import os
import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytesseract
from flask import Flask
from docx import Document
from docx.oxml import OxmlElement

from routes.candidates import candidates
from utils.resume_parser import (
    extract_email,
    extract_name,
    extract_name_from_filename,
    extract_text_from_docx,
    extract_text_from_pdf,
)
from utils.resume_converter import convert_resume_to_pdf


class FakeCursor:
    def __init__(self):
        self.executed = []
        self.results = iter([(1, "Data Engineer"), ("Python", "SQL"), (7,)])
        self.parameters = []

    def execute(self, query, parameters=None):
        self.executed.append(query)
        self.parameters.append(parameters)

    def fetchone(self):
        return next(self.results)


class FakeConnection:
    def __init__(self):
        self.fake_cursor = FakeCursor()

    def cursor(self):
        return self.fake_cursor

    def commit(self):
        pass

    def close(self):
        pass


class ResumeTextExtractionTests(unittest.TestCase):
    def test_filename_with_underscore_is_split_into_name_tokens(self):
        self.assertEqual(
            extract_name_from_filename("sreenivas_palika.pdf"),
            "Sreenivas Palika",
        )
        self.assertEqual(
            extract_name("Resume starts here", "sreenivas_palika.pdf"),
            "Sreenivas Palika",
        )

    def test_concatenated_filename_name_falls_back_to_resume_text(self):
        self.assertEqual(
            extract_name_from_filename("devaraboinaaravind.pdf"),
            "",
        )
        self.assertEqual(
            extract_name(
                "Devaraboina Aravind\nSenior Data Engineer\nSkills: Python",
                "devaraboinaaravind.pdf",
            ),
            "Devaraboina Aravind",
        )

    def test_single_token_filename_does_not_become_full_name(self):
        self.assertEqual(extract_name("", "sreenivas.pdf"), "")

    def test_name_labels_inside_client_and_project_headings_are_ignored(self):
        text = (
            "Client Name: Altice USA\n"
            "Aravind Devaraboina\n"
            "Senior Data Engineer"
        )
        self.assertEqual(
            extract_name(text, "Aravind_Devaraboina_CV.docx"),
            "Aravind Devaraboina",
        )

        text = "Project Name: IT Systems Consolidation\nSREENIVASA RAO PALIKA\n"
        self.assertEqual(
            extract_name(text, "Sreenivas_Palika_Senior_Data_Engineer.docx"),
            "Sreenivasa Rao Palika",
        )

    def test_single_word_resume_name_is_used_when_filename_prefix_matches(self):
        self.assertEqual(
            extract_name(
                "Ankesh\nData Engineer\nSpecialized in cloud data engineering",
                "Ankesh_Data Engineer.pdf",
            ),
            "Ankesh",
        )

    def test_names_in_formatted_resume_headers_are_extracted(self):
        cases = [
            ("ASHISH | PATIL\nDATA ENGINEERING | SPECIALIST", "Ashish Patil"),
            ("VAIBHAV | NIRKHEY\n☎ +91 9763259519", "Vaibhav Nirkhey"),
            ("Mansi | Goel\nPower BI Developer", "Mansi Goel"),
            (
                "Jangam Suresh Babu Email: jangamsuresh05@gmail.com",
                "Jangam Suresh Babu",
            ),
        ]

        for text, expected_name in cases:
            with self.subTest(expected_name=expected_name):
                self.assertEqual(extract_name(text), expected_name)

    def _mock_pdf(self, page):
        pdf = MagicMock()
        pdf.__enter__.return_value.pages = [page]
        return pdf

    @patch("utils.resume_parser.pdfplumber.open")
    def test_pdf_extraction_preserves_columns(self, open_pdf):
        layout = (
            "Jane Example\n"
            "Experience                              Skills\n"
            "Senior Data Engineer                    Python SQL Airflow"
        )
        page = MagicMock()
        page.extract_text.return_value = layout
        page.images = []
        open_pdf.return_value = self._mock_pdf(page)

        text = extract_text_from_pdf("two-column.pdf")

        self.assertEqual(
            text,
            "Jane Example\nExperience | Skills\n"
            "Senior Data Engineer | Python SQL Airflow",
        )
        self.assertEqual(extract_name(text), "Jane Example")
        page.extract_text.assert_called_once_with(layout=True)
        page.to_image.assert_not_called()

    @patch("utils.resume_parser.pytesseract.image_to_string")
    @patch("utils.resume_parser.pdfplumber.open")
    def test_image_only_pdf_uses_ocr(self, open_pdf, image_to_string):
        page = MagicMock()
        page.extract_text.return_value = None
        page.images = [{"name": "scanned-page"}]
        page.to_image.return_value = SimpleNamespace(original=object())
        open_pdf.return_value = self._mock_pdf(page)
        image_to_string.return_value = "Morgan Lee morgan@example.com"

        text = extract_text_from_pdf("scanned.pdf")

        self.assertEqual(text, "Morgan Lee morgan@example.com")
        page.to_image.assert_called_once_with(resolution=300)
        image_to_string.assert_called_once_with(page.to_image.return_value.original)

    @patch("utils.resume_parser.pytesseract.image_to_string")
    @patch("utils.resume_parser.pdfplumber.open")
    def test_sparse_text_on_image_page_is_supplemented_with_ocr(
        self, open_pdf, image_to_string
    ):
        page = MagicMock()
        page.extract_text.return_value = "Morgan Lee"
        page.images = [{"name": "resume-image"}]
        page.to_image.return_value = SimpleNamespace(original=object())
        open_pdf.return_value = self._mock_pdf(page)
        image_to_string.return_value = "Morgan Lee Experience Python SQL"

        text = extract_text_from_pdf("photo-resume.pdf")

        self.assertIn("Morgan Lee", text)
        self.assertIn("Experience Python SQL", text)
        image_to_string.assert_called_once_with(page.to_image.return_value.original)

    @patch("utils.resume_parser.pytesseract.image_to_string")
    @patch("utils.resume_parser.pdfplumber.open")
    def test_image_only_pdf_explains_missing_tesseract(
        self, open_pdf, image_to_string
    ):
        page = MagicMock()
        page.extract_text.return_value = None
        page.images = [{"name": "scanned-page"}]
        open_pdf.return_value = self._mock_pdf(page)
        image_to_string.side_effect = pytesseract.TesseractNotFoundError()

        with self.assertRaisesRegex(RuntimeError, "Install Tesseract OCR"):
            extract_text_from_pdf("scanned.pdf")

    def test_docx_extraction_includes_table_columns(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "two-column.docx"
            document = Document()
            document.add_paragraph("Morgan Lee")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Experience"
            table.cell(0, 1).text = "Python SQL"
            document.save(path)

            text = extract_text_from_docx(str(path))

        self.assertIn("Morgan Lee", text)
        self.assertIn("Experience | Python SQL", text)

    def test_docx_extraction_includes_textbox_email(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "textbox-email.docx"
            document = Document()
            paragraph = document.add_paragraph()
            textbox = OxmlElement("w:txbxContent")
            textbox_paragraph = OxmlElement("w:p")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "scsnhd0421@gmail.com"
            run.append(text)
            textbox_paragraph.append(run)
            textbox.append(textbox_paragraph)
            paragraph._p.append(textbox)
            document.save(path)

            extracted = extract_text_from_docx(str(path))

        self.assertIn("scsnhd0421@gmail.com", extracted)
        self.assertEqual(extract_email(extracted), "scsnhd0421@gmail.com")


class ResumeConversionTests(unittest.TestCase):
    @patch("utils.resume_converter.subprocess.run")
    @patch("utils.resume_converter._find_libreoffice", return_value="soffice")
    def test_docx_is_converted_to_pdf_with_headless_libreoffice(
        self, find_libreoffice, run
    ):
        def create_converted_pdf(command, **kwargs):
            self.assertTrue(kwargs["check"])
            source_path = command[-1]
            output_directory = command[command.index("--outdir") + 1]
            output_path = os.path.join(
                output_directory,
                f"{Path(source_path).stem}.pdf",
            )
            Path(output_path).write_bytes(b"%PDF-converted")
            return SimpleNamespace(returncode=0, stdout="", stderr="")

        run.side_effect = create_converted_pdf

        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "resume.docx"
            source_path.write_bytes(b"docx content")

            pdf_path = convert_resume_to_pdf(str(source_path), temp_dir)
            converted_bytes = Path(pdf_path).read_bytes()

        self.assertEqual(Path(pdf_path).suffix, ".pdf")
        self.assertEqual(converted_bytes, b"%PDF-converted")
        command = run.call_args.args[0]
        self.assertIn("--headless", command)
        self.assertEqual(command[command.index("--convert-to") + 1], "pdf")
        self.assertTrue(run.call_args.kwargs["check"])
        find_libreoffice.assert_called_once()


    @patch("routes.candidates.parse_resume")
    @patch("routes.candidates.get_connection")
    @patch("routes.candidates.ensure_schema")
    def test_failed_parse_is_not_inserted_as_a_candidate(
        self, ensure_schema, get_connection, parse_resume
    ):
        connection = FakeConnection()
        get_connection.return_value = connection
        parse_resume.return_value = {
            "status": "failed",
            "error": "OCR is not configured",
        }
        app = Flask(__name__)
        app.register_blueprint(candidates)

        with tempfile.TemporaryDirectory() as temp_dir:
            with patch("routes.candidates.os.getcwd", return_value=temp_dir):
                response = app.test_client().post(
                    "/api/candidates/upload",
                    data={
                        "job_id": "1",
                        "resumes": (io.BytesIO(b"resume"), "scanned.pdf"),
                    },
                    content_type="multipart/form-data",
                )

        self.assertEqual(response.status_code, 201)
        result = response.get_json()
        self.assertEqual(result["created_count"], 0)
        self.assertEqual(result["failed"][0]["fileName"], "scanned.pdf")
        ensure_schema.assert_called_once()
        self.assertEqual(len(connection.fake_cursor.executed), 2)
        self.assertFalse(
            any(
                "INSERT INTO dbo.Candidates" in query
                for query in connection.fake_cursor.executed
            )
        )

    @patch("routes.candidates.parse_resume")
    @patch("routes.candidates.get_connection")
    @patch("routes.candidates.ensure_schema")
    def test_uploaded_resume_is_linked_to_created_candidate(
        self, ensure_schema, get_connection, parse_resume
    ):
        connection = FakeConnection()
        get_connection.return_value = connection
        parse_resume.return_value = {
            "status": "success",
            "full_name": "Morgan Lee",
            "email": "morgan@example.com",
            "phone": None,
            "current_role": "Engineer",
            "location": "Hyderabad",
            "experience_years": 6,
            "current_ctc": None,
            "notice_period": None,
            "skills": ["Python", "SQL"],
            "text": "Morgan Lee",
        }
        app = Flask(__name__)
        app.register_blueprint(candidates)

        with tempfile.TemporaryDirectory() as temp_dir:
            with patch("routes.candidates.os.getcwd", return_value=temp_dir):
                response = app.test_client().post(
                    "/api/candidates/upload",
                    data={
                        "job_id": "1",
                        "resumes": (io.BytesIO(b"resume"), "morgan_lee.pdf"),
                    },
                    content_type="multipart/form-data",
                )

            self.assertEqual(response.status_code, 201)
            self.assertEqual(
                response.get_json()["pdf_ready_count"], 1
            )
            self.assertEqual(response.get_json()["converted_count"], 0)
            self.assertEqual(response.get_json()["parsed_count"], 1)
            insert_query, insert_parameters = list(
                zip(
                    connection.fake_cursor.executed,
                    connection.fake_cursor.parameters,
                )
            )[-1]
            self.assertIn("resume_storage_name", insert_query)
            self.assertIn("resume_original_name", insert_query)
            self.assertEqual(insert_parameters[-1], "morgan_lee.pdf")
            self.assertTrue(
                os.path.isfile(
                    os.path.join(
                        temp_dir,
                        "uploads",
                        "resumes",
                        insert_parameters[-2],
                    )
                )
            )
            ensure_schema.assert_called_once()
            ensure_schema.assert_called_once()

    @patch("routes.candidates.convert_resume_to_pdf")
    @patch("routes.candidates.parse_resume")
    @patch("routes.candidates.get_connection")
    @patch("routes.candidates.ensure_schema")
    def test_uploaded_docx_is_converted_before_parsing_and_stored_as_pdf(
        self, ensure_schema, get_connection, parse_resume, convert_to_pdf
    ):
        connection = FakeConnection()
        get_connection.return_value = connection
        parse_resume.return_value = {
            "status": "success",
            "full_name": "Morgan Lee",
            "email": "morgan@example.com",
            "phone": None,
            "current_role": "Engineer",
            "location": "Hyderabad",
            "experience_years": 6,
            "current_ctc": None,
            "notice_period": None,
            "skills": ["Python", "SQL"],
            "text": "Morgan Lee",
        }
        app = Flask(__name__)
        app.register_blueprint(candidates)

        def create_converted_pdf(source_path, output_directory):
            pdf_path = os.path.join(
                output_directory, f"{Path(source_path).stem}.pdf"
            )
            Path(pdf_path).write_bytes(b"%PDF-converted")
            return pdf_path

        convert_to_pdf.side_effect = create_converted_pdf

        with tempfile.TemporaryDirectory() as temp_dir:
            with patch("routes.candidates.os.getcwd", return_value=temp_dir):
                response = app.test_client().post(
                    "/api/candidates/upload",
                    data={
                        "job_id": "1",
                        "resumes": (io.BytesIO(b"docx content"), "morgan_lee.docx"),
                    },
                    content_type="multipart/form-data",
                )

            self.assertEqual(response.status_code, 201)
            self.assertEqual(response.get_json()["created_count"], 1)
            self.assertEqual(response.get_json()["pdf_ready_count"], 1)
            self.assertEqual(response.get_json()["converted_count"], 1)
            self.assertEqual(response.get_json()["parsed_count"], 1)

            convert_to_pdf.assert_called_once()
            parsed_pdf_path, parsed_file_name = parse_resume.call_args.args
            self.assertTrue(parsed_pdf_path.endswith(".pdf"))
            self.assertEqual(parsed_file_name, "morgan_lee.pdf")
            self.assertEqual(Path(parsed_pdf_path).read_bytes(), b"%PDF-converted")

            insert_query, insert_parameters = list(
                zip(
                    connection.fake_cursor.executed,
                    connection.fake_cursor.parameters,
                )
            )[-1]
            self.assertIn("resume_storage_name", insert_query)
            self.assertTrue(insert_parameters[-2].endswith(".pdf"))
            self.assertEqual(insert_parameters[-1], "morgan_lee.pdf")
            self.assertTrue(
                os.path.isfile(
                    os.path.join(
                        temp_dir,
                        "uploads",
                        "resumes",
                        insert_parameters[-2],
                    )
                )
            )
            ensure_schema.assert_called_once()


if __name__ == "__main__":
    unittest.main()
