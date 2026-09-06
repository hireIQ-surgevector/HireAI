"""
text_extractor.py
Shared text extraction for resumes (PDF and DOCX), with OCR fallback for
scanned/image-based pages. Used by both regex_fields.py and llm_fields.py
so text is only extracted once per resume.

Requirements:
    pip install pymupdf python-docx pytesseract pillow --break-system-packages
    Tesseract OCR binary installed separately.
"""

import io
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from PIL import Image
import pytesseract

MIN_TEXT_CHARS_PER_PAGE = 40


def extract_text_from_pdf(path: Path) -> str:
    doc = fitz.open(path)
    pages_text = []
    for page in doc:
        text = page.get_text("text").strip()
        if len(text) < MIN_TEXT_CHARS_PER_PAGE:
            pix = page.get_pixmap(dpi=300)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            text = pytesseract.image_to_string(img)
        pages_text.append(text)
    doc.close()
    return "\n".join(pages_text)


def extract_text_from_docx(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    elif suffix in (".docx", ".doc"):
        return extract_text_from_docx(path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")